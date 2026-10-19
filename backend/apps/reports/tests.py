import csv
import io
from datetime import date, timedelta
from decimal import Decimal
from unittest import mock

from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.accounts.models import PermissionGrant, Role
from apps.accounts.tests import login
from apps.audit.models import AuditLog
from apps.billing import services as bs
from apps.billing.tests import new_customer
from apps.core.testing import make_office, make_package, make_user, make_village
from apps.customers import services as cs
from apps.customers.models import Connection, Customer
from apps.expenses import services as ex
from apps.finance import services as fs
from apps.inventory import services as inv
from apps.operations import services as ops
from apps.reports import export
from apps.reports.registry import REPORTS


class ExportWriterTests(TestCase):
    def test_formula_injection_is_neutralised_in_every_format(self):
        evil = ['=HYPERLINK("http://evil","x")', "+1+1", "-2+3", "@SUM(A1)", "\t=1", "\r=1", "normal", "ok=fine"]
        rows = [[v, 5, Decimal("-12.50")] for v in evil]
        reader = list(csv.reader(io.StringIO(export.to_csv(["a", "b", "c"], rows).decode("utf-8-sig"))))
        for original, row in zip(evil, reader[1:]):
            if original[:1] in ("=", "+", "-", "@", "\t", "\r"):
                self.assertTrue(row[0].startswith("'"), original)
            else:
                self.assertEqual(row[0], original)
            self.assertEqual(row[2], "-12.50")           # real negative numbers are left alone
        ws = load_workbook(io.BytesIO(export.to_xlsx("T", ["a", "b", "c"], rows))).active
        self.assertEqual(ws["A2"].value, "'" + evil[0])
        self.assertNotEqual(ws["A2"].data_type, "f")      # never stored as a formula
        self.assertEqual(ws["C2"].value, -12.5)
        self.assertTrue(export.to_pdf("T", ["a"], [[evil[0]], ["<b>&"]]).startswith(b"%PDF"))

    def test_unicode_survives_csv(self):
        out = export.to_csv(["نام"], [["علی خان"]]).decode("utf-8-sig")
        self.assertIn("علی خان", out)


class ReportBase(TestCase):
    def setUp(self):
        self.a, self.b = make_office("AAA"), make_office("BBB")
        self.va, self.vb = make_village(self.a, "KHT"), make_village(self.b, "BDR")
        self.owner = make_user(Role.OWNER)
        self.mgr = make_user(Role.MANAGER, [self.a])
        self.mgr_b = make_user(Role.MANAGER, [self.b])
        self.staff = make_user(Role.OFFICE_STAFF, [self.a])
        self.col = make_user(Role.COLLECTOR, [self.a])
        self.tech = make_user(Role.TECHNICIAN, [self.a])
        self.pkg = make_package()

    def api(self, u):
        c = APIClient(); login(c, u); return c

    def cust(self, name="Ali", cnic="3520212345671", staff=None, village=None, office=None):
        return cs.register_customer(actor=staff or self.staff, confirm_duplicates=True, data={
            "full_name": name, "phone": "0300", "cnic": cnic, "permanent_address": "x", "village": village or self.va, "office": office or self.a})

    def conn(self, customer):
        return cs.create_connection(actor=self.staff, customer=customer, package=self.pkg, installation_address="x")

    def pay(self, conn, actor=None, method="cash", ref=""):
        inv_ = bs.create_invoice(actor=self.staff, connection=conn)[0]
        return bs.record_payment(actor=actor or self.col, invoice=inv_, method=method, amount=inv_.total_amount, reference=ref)

    def get(self, user, name, **q):
        return self.api(user).get(f"/api/v1/reports/{name}/", q)


class ReportTests(ReportBase):
    def test_csv_export_escapes_formulas_and_audits_the_export(self):
        self.cust(name='=cmd|"/c calc"!A1')
        self.cust(name="-1+1")
        r = self.get(self.mgr, "customers", fmt="csv")
        self.assertEqual((r.status_code, r["Content-Type"]), (200, "text/csv; charset=utf-8"))
        self.assertIn("attachment", r["Content-Disposition"])
        names = [row[1] for row in csv.reader(io.StringIO(r.content.decode("utf-8-sig")))][1:]
        self.assertEqual(sorted(names), sorted(["'=cmd|\"/c calc\"!A1", "'-1+1"]))
        log = AuditLog.objects.get(action="report.exported")
        self.assertEqual((log.after["format"], log.after["rows"]), ("csv", 2))

    def test_cnic_is_masked_unless_the_user_may_see_it(self):
        self.cust()
        row = lambda u: list(csv.reader(io.StringIO(self.get(u, "customers", fmt="csv").content.decode("utf-8-sig"))))[1][3]
        self.assertEqual(row(self.mgr), "*****-*******-1")
        PermissionGrant.objects.create(employee=self.mgr, office=self.a, code="customers.view_cnic")
        self.assertEqual(row(self.mgr), "35202-1234567-1")
        self.assertEqual(self.get(self.mgr, "customers").json()["results"][0][3], "35202-1234567-1")   # preview follows the same rule

    def test_office_scope_applies_to_every_report(self):
        ca, cb = self.cust("Ali"), self.cust("Bashir", staff=make_user(Role.OFFICE_STAFF, [self.b]), village=self.vb, office=self.b)
        conn = self.conn(ca)
        self.pay(conn)
        self.assertEqual(self.get(self.mgr, "customers").json()["count"], 1)
        self.assertEqual(self.get(self.mgr_b, "customers").json()["count"], 1)
        self.assertEqual(self.get(self.owner, "customers").json()["count"], 2)
        self.assertEqual(self.get(self.mgr_b, "collections").json()["count"], 0)
        self.assertEqual(self.get(self.mgr_b, "invoices").json()["count"], 0)
        self.assertEqual(self.get(self.mgr, "collections", office=self.b.pk).json()["count"], 0)    # cannot widen scope with a filter

    def test_export_needs_its_own_permission_but_preview_does_not(self):
        self.cust()
        self.assertEqual(self.get(self.staff, "customers").status_code, 200)               # staff may preview...
        self.assertEqual(self.get(self.staff, "customers", fmt="csv").status_code, 403)    # ...but not export
        PermissionGrant.objects.create(employee=self.staff, office=self.a, code="reports.export")
        self.assertEqual(self.get(self.staff, "customers", fmt="csv").status_code, 403)    # customer data also needs customers.export
        PermissionGrant.objects.create(employee=self.staff, office=self.a, code="customers.export")
        self.assertEqual(self.get(self.staff, "customers", fmt="csv").status_code, 200)
        self.assertEqual(self.get(self.tech, "customers").status_code, 403)
        self.assertEqual(self.get(self.mgr, "nope").status_code, 404)
        names = {r["name"] for r in self.api(self.tech).get("/api/v1/reports/").json()}
        self.assertNotIn("customers", names)
        self.assertNotIn("collections", names)

    def test_xlsx_and_pdf(self):
        conn = self.conn(self.cust())
        self.pay(conn)
        x = self.get(self.mgr, "collections", fmt="xlsx")
        self.assertEqual(x["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        ws = load_workbook(io.BytesIO(x.content)).active
        self.assertEqual([c.value for c in ws[1]][:2], ["Confirmed", "Office"])
        self.assertEqual(float(ws["I2"].value), 1500.0)   # Amount column
        p = self.get(self.mgr, "collections", fmt="pdf")
        self.assertEqual((p["Content-Type"], p.content[:4]), ("application/pdf", b"%PDF"))
        self.assertEqual(self.get(self.mgr, "collections", fmt="docx").status_code, 400)

    def test_filters_dates_status_and_method(self):
        c1, c2 = self.conn(self.cust("A")), self.conn(self.cust("B"))
        self.pay(c1)
        self.pay(c2, actor=self.staff, method="easypaisa", ref="E1")
        self.assertEqual(self.get(self.mgr, "collections").json()["count"], 1)                     # unverified digital excluded
        self.assertEqual(self.get(self.mgr, "pending-verification").json()["count"], 1)
        self.assertEqual(self.get(self.mgr, "collections", method="easypaisa").json()["count"], 0)
        today = str(timezone.localdate())
        self.assertEqual(self.get(self.mgr, "collections", date_from=today, date_to=today).json()["count"], 1)
        tomorrow = str(timezone.localdate() + timedelta(days=1))
        self.assertEqual(self.get(self.mgr, "collections", date_from=tomorrow).json()["count"], 0)
        self.assertEqual(self.get(self.mgr, "collections", date_from="31-12-2026").status_code, 400)
        self.assertEqual(self.get(self.mgr, "outstanding").json()["count"], 1)                     # invoice with pending payment
        self.assertEqual(self.get(self.mgr, "invoices", status="paid").json()["count"], 1)

    def test_preview_is_paginated_server_side(self):
        for i in range(7):
            self.cust(f"Cust {i}")
        r = self.get(self.mgr, "customers", page=2, page_size=3).json()
        self.assertEqual((r["count"], r["page"], len(r["results"])), (7, 2, 3))
        self.assertEqual(len(self.get(self.mgr, "customers", page=3, page_size=3).json()["results"]), 1)

    def test_row_cap_protects_serverless_limits(self):
        for i in range(3):
            self.cust(f"C{i}")
        with mock.patch("apps.reports.views.MAX_ROWS", 2):
            r = self.get(self.mgr, "customers", fmt="csv")
        self.assertEqual(r.status_code, 400)
        self.assertIn("Narrow the filters", r.json()["error"]["message"])
        with mock.patch("apps.reports.views.MAX_BYTES", 10):
            self.assertEqual(self.get(self.mgr, "customers", fmt="csv").status_code, 400)

    def test_other_reports_run(self):
        c = self.cust(); conn = self.conn(c)
        ops.create_complaint(actor=self.staff, customer=c, category="other", description="x")
        ex.create_expense(actor=self.staff, office=self.a, category="fuel", description="d", amount="100", expense_date=timezone.localdate(), payment_method="cash")
        item = inv.create_item(actor=self.mgr, office=self.a, name="Cable", category="cable", minimum_stock=5)
        inv.record_purchase(actor=self.mgr, item=item, quantity="3", unit_price="1", purchase_date=timezone.localdate())
        Connection.objects.filter(pk=conn.pk).update(status="active", expires_at=timezone.now() + timedelta(days=2))
        for name, n in (("connections", 1), ("expiry", 1), ("complaints", 1), ("expenses", 1), ("inventory", 1), ("stock-movements", 1),
                        ("technician-workload", 0), ("attendance", 0), ("leave", 0), ("refunds", 0), ("cash-reconciliation", 0)):
            r = self.get(self.mgr, name)
            self.assertEqual((r.status_code, r.json()["count"]), (200, n), name)
        self.assertEqual(self.get(self.mgr, "inventory").json()["results"][0][-1], "Yes")   # 3 in stock, minimum 5 -> low
        self.assertEqual(self.get(self.mgr, "expiry", until=str(timezone.localdate() - timedelta(days=5))).json()["count"], 0)
        self.assertEqual(set(REPORTS), {r["name"] for r in self.api(self.owner).get("/api/v1/reports/").json()})


class DashboardTests(ReportBase):
    def test_figures_match_the_underlying_records_and_keep_categories_apart(self):
        c1, c2, c3 = self.conn(self.cust("A")), self.conn(self.cust("B")), self.conn(self.cust("C"))
        self.pay(c1)                                                         # cash 1500 confirmed
        pend = self.pay(c2, actor=self.staff, method="jazzcash", ref="J1")   # digital 1500 pending
        rev = self.pay(c3)                                                   # cash 1500, then reversed
        corr = fs.request_correction(actor=self.staff, payment=rev, reason="wrong")
        fs.approve_correction(actor=self.mgr, correction=corr)
        cs.complete_installation(actor=self.mgr, connection=Connection.objects.get(pk=c1.pk))
        d = self.api(self.mgr).get("/api/v1/dashboard/").json()
        self.assertEqual(d["customers"], {"total": 3, "active": 3})
        self.assertEqual(d["connections"]["by_status"], {"pending_installation": 2, "pending_approval": 1})
        self.assertEqual((d["connections"]["pending_installations"], d["connections"]["pending_approvals"]), (2, 1))
        col = d["collections"]
        self.assertEqual((col["today"]["cash"], col["today"]["digital_verified"]), ("1500.00", "0.00"))
        self.assertEqual((col["pending_verification"]["count"], col["pending_verification"]["amount"]), (1, "1500.00"))
        self.assertEqual((col["reversed_this_month"], col["undeposited_cash"]), ("1500.00", "1500.00"))
        self.assertEqual(d["outstanding_invoices"], {"count": 2, "amount": "3000.00"})   # c2 (pending) + c3 (reopened)
        bs.verify_payment(actor=self.mgr, payment=pend)
        d2 = self.api(self.mgr).get("/api/v1/dashboard/").json()
        self.assertEqual((d2["collections"]["today"]["digital_verified"], d2["collections"]["pending_verification"]["count"]), ("1500.00", 0))
        self.assertIn("not profit", d["scope_note"])

    def test_sections_follow_permissions_and_office_scope(self):
        self.pay(self.conn(self.cust()))
        tech = self.api(self.tech).get("/api/v1/dashboard/").json()
        for hidden in ("customers", "collections", "expenses", "outstanding_invoices", "office_income_month"):
            self.assertNotIn(hidden, tech, hidden)
        self.assertIn("tasks", tech)
        self.assertEqual(self.api(self.mgr_b).get("/api/v1/dashboard/").json()["collections"]["month"]["cash"], "0.00")
        owner = self.api(self.owner).get("/api/v1/dashboard/").json()
        self.assertEqual(owner["office_income_month"], {"AAA": "1500.00"})
        self.assertEqual(self.api(self.owner).get(f"/api/v1/dashboard/?office={self.b.pk}").json()["customers"]["total"], 0)
        col = self.api(self.col).get("/api/v1/dashboard/").json()
        self.assertEqual(col["collections"]["today"]["cash"], "1500.00")
        self.assertNotIn("expenses", col)

    def test_complaints_inventory_attendance_and_leave_sections(self):
        c = self.cust()
        k = ops.create_complaint(actor=self.staff, customer=c, category="other", description="x", priority="urgent")
        ops.create_complaint(actor=self.staff, customer=c, category="other", description="y")
        ops.assign(actor=self.staff, complaint=k, technician=self.tech)
        item = inv.create_item(actor=self.mgr, office=self.a, name="Cable", category="cable", minimum_stock=50)
        from apps.hr import services as hr
        hr.check_in(actor=self.staff)
        hr.request_leave(actor=self.staff, start_date=timezone.localdate() + timedelta(days=3), end_date=timezone.localdate() + timedelta(days=4), reason="x")
        ex.create_expense(actor=self.staff, office=self.a, category="fuel", description="d", amount="250", expense_date=timezone.localdate(), payment_method="cash")
        d = self.api(self.mgr).get("/api/v1/dashboard/").json()
        self.assertEqual(d["complaints"], {"open": 2, "by_priority": {"urgent": 1, "medium": 1}, "unassigned": 1})
        self.assertEqual(d["inventory"]["low_stock_items"], 1)
        self.assertEqual((d["attendance_today"]["checked_in"], d["leave_pending"]), (1, 1))
        self.assertEqual((d["expenses"]["pending_approval_count"], d["expenses"]["pending_approval_amount"], d["expenses"]["approved_this_month"]), (1, "250.00", "0.00"))
        t = self.api(self.tech).get("/api/v1/dashboard/").json()
        self.assertEqual(t["complaints"]["open"], 1)         # technicians see only their own work
        self.assertFalse(t["my_attendance"]["checked_in"])


class AuditLogApiTests(ReportBase):
    def test_only_authorised_users_read_audit_logs(self):
        self.cust()
        self.assertEqual(self.api(self.mgr).get("/api/v1/audit-logs/").status_code, 403)
        self.assertEqual(self.api(self.staff).get("/api/v1/audit-logs/").status_code, 403)
        PermissionGrant.objects.create(employee=self.mgr, office=None, code="audit.view")
        self.assertEqual(self.api(self.mgr).get("/api/v1/audit-logs/").status_code, 200)
        o = self.api(self.owner)
        r = o.get("/api/v1/audit-logs/?action=customer.created").json()
        self.assertEqual(r["count"], 1)
        self.assertEqual((r["results"][0]["entity_type"], r["results"][0]["actor_label"]), ("customers.customer", self.staff.username))
        self.assertNotIn("3520212345671", str(r))                          # CNIC masked in the log
        self.assertEqual(o.get("/api/v1/audit-logs/?entity_type=customers.customer&date_from=2000-01-01").json()["count"], 1)
        self.assertEqual(o.get("/api/v1/audit-logs/?action_prefix=customer.").json()["count"], 1)

    def test_audit_api_is_read_only(self):
        o = self.api(self.owner)
        entry = AuditLog.objects.first() or self.cust() and AuditLog.objects.first()
        for method in ("post", "patch", "delete", "put"):
            self.assertIn(getattr(o, method)(f"/api/v1/audit-logs/{entry.pk}/", {}, format="json").status_code, (403, 405), method)
        self.assertEqual(o.get(f"/api/v1/audit-logs/{entry.pk}/").status_code, 200)
