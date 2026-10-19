"""File writers. Every cell that comes from user-controlled text is neutralised against spreadsheet formula injection."""
import csv
import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n")


def safe_cell(value):
    """Numbers stay numbers; text starting with a formula character gets a leading apostrophe."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (int, float, Decimal)):
        return value
    text = str(value)
    return "'" + text if text.startswith(FORMULA_PREFIXES) else text


def _text(value):
    if isinstance(value, Decimal) and not isinstance(value, bool):
        return str(value.quantize(Decimal("0.01")))
    return str(safe_cell(value))


def to_csv(columns, rows) -> bytes:
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow([_text(c) for c in columns])
    for r in rows:
        w.writerow([_text(c) for c in r])
    return ("\ufeff" + out.getvalue()).encode("utf-8")  # BOM so Excel reads UTF-8 (Urdu names) correctly


def to_xlsx(title, columns, rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = (title or "Report")[:31]
    ws.append([str(c) for c in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([safe_cell(c) for c in r])
    for i, col in enumerate(columns, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(40, max(12, len(str(col)) + 4))
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(title, columns, rows, subtitle="") -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=28, bottomMargin=28, title=title)
    styles = getSampleStyleSheet()
    cell = styles["BodyText"]
    cell.fontSize = 7
    cell.leading = 9
    data = [[Paragraph(f"<b>{_esc(c)}</b>", cell) for c in columns]]
    for r in rows:
        data.append([Paragraph(_esc(_text(c)), cell) for c in r])
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.25, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef2ff")),
                               ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story = [Paragraph(_esc(title), styles["Title"])]
    if subtitle:
        story += [Paragraph(_esc(subtitle), styles["Normal"]), Spacer(1, 6)]
    story.append(table)
    doc.build(story)
    return buf.getvalue()


def _esc(text):
    return str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
